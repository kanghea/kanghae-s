#!/usr/bin/env python3
"""PPTX → 갤러리 미리보기 이미지 + 잠금 슬라이드 베일 데이터.

공개 저장소·배포본에는 미리보기 슬라이드(기본 1~3장)만 원본 해상도로 들어간다.
나머지 슬라이드는 8px 폭 색 견본(data URI)만 남겨 "가려진 작품" 효과에 쓴다 —
글자·얼굴을 알아볼 수 없는 크기라 내용이 새지 않는다. PPTX 원본은 커밋하지 않는다.

사용:
    python3 scripts/import_deck.py ~/Downloads/deck.pptx --slug my-deck
    python3 scripts/import_deck.py deck.pptx --slug my-deck --preview 1,2,5

산출물:
    public/gallery/<slug>/slide-NN.webp      미리보기 슬라이드(폭 --width, 기본 1920)
    src/content/gallery/<slug>.deck.json     장수·비율·미리보기 경로·베일 data URI

필요 도구: python-pptx, Pillow, LibreOffice(soffice, Impress 포함), poppler(pdftoppm).
슬라이드가 전부 전면 이미지 1장인 덱(이미지로 내보낸 PPT)은 LibreOffice 없이 원본
이미지를 그대로 뽑는다. 그 밖의 덱은 LibreOffice 로 렌더하므로 덱이 쓰는 글꼴이
설치돼 있어야 원본과 같다 — 없는 글꼴은 실행 시 경고로 알린다.
"""

from __future__ import annotations

import argparse
import base64
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent.parent
VEIL_WIDTH = 8  # 잠금 슬라이드 색 견본 폭(px). 키우면 내용이 비칠 수 있다.
SLUG_RE = re.compile(r"^[a-z0-9]+(?:-[a-z0-9]+)*$")


def parse_preview(value: str, count: int) -> list[int]:
    slides = sorted({int(v) for v in value.split(",") if v.strip()})
    bad = [s for s in slides if s < 1 or s > count]
    if bad:
        sys.exit(f"--preview 범위 밖 슬라이드: {bad} (총 {count}장)")
    return slides


def full_bleed_images(prs: Presentation, preview: list[int]) -> list[bytes] | None:
    """모든 슬라이드가 슬라이드 크기의 그림 1장이면 그 원본 바이트 목록.

    잠금 슬라이드에 붙은 동영상(MEDIA)은 베일 색에 영향이 없어 허용한다.
    미리보기 슬라이드에 다른 도형이 있으면 렌더 결과가 달라지므로 None.
    """
    out: list[bytes] = []
    w, h = prs.slide_width, prs.slide_height
    for n, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        pics = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        others = [s for s in shapes if s.shape_type != MSO_SHAPE_TYPE.PICTURE]
        if len(pics) != 1:
            return None
        if others and (n in preview or any(s.shape_type != MSO_SHAPE_TYPE.MEDIA for s in others)):
            return None
        pic = pics[0]
        if abs(pic.left) > w * 0.01 or abs(pic.top) > h * 0.01:
            return None
        if abs(pic.width - w) > w * 0.01 or abs(pic.height - h) > h * 0.01:
            return None
        out.append(pic.image.blob)
    return out


def deck_fonts(prs: Presentation) -> set[str]:
    fonts: set[str] = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        fonts.add(run.font.name)
    return fonts


def warn_missing_fonts(fonts: set[str]) -> None:
    if not shutil.which("fc-match"):
        return
    for name in sorted(fonts):
        got = subprocess.run(
            ["fc-match", "-f", "%{family}", name], capture_output=True, text=True
        ).stdout
        if name.lower() not in got.lower():
            print(f"경고: 글꼴 '{name}' 없음 → '{got}' 로 대체 렌더됩니다.", file=sys.stderr)


def repack(src: Path, dst: Path) -> None:
    """[Content_Types].xml 을 맨 앞에 두도록 다시 묶는다(일부 변환기 산출물 호환)."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        names = [n for n in zin.namelist() if not n.endswith("/")]
        names.sort(key=lambda n: n != "[Content_Types].xml")
        for n in names:
            zout.writestr(n, zin.read(n))


def render_with_libreoffice(pptx: Path, tmp: Path, preview: list[int], width: int) -> tuple[dict[int, Image.Image], list[Image.Image]]:
    for tool in ("soffice", "pdftoppm"):
        if not shutil.which(tool):
            sys.exit(f"{tool} 가 필요합니다(LibreOffice Impress · poppler-utils).")
    work = tmp / "deck.pptx"
    repack(pptx, work)
    profile = tmp / "lo-profile"
    subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
         "--convert-to", "pdf", "--outdir", str(tmp), str(work)],
        check=True, capture_output=True,
    )
    pdf = tmp / "deck.pdf"
    if not pdf.exists():
        sys.exit("LibreOffice 변환 실패 — libreoffice-impress 가 설치돼 있는지 확인하세요.")

    previews: dict[int, Image.Image] = {}
    for n in preview:
        prefix = tmp / f"p{n}"
        subprocess.run(
            ["pdftoppm", "-png", "-f", str(n), "-l", str(n), "-scale-to-x", str(width),
             "-scale-to-y", "-1", "-singlefile", str(pdf), str(prefix)],
            check=True,
        )
        previews[n] = Image.open(f"{prefix}.png").convert("RGB")

    small = tmp / "small"
    subprocess.run(["pdftoppm", "-png", "-scale-to-x", "64", "-scale-to-y", "-1", str(pdf), str(small)], check=True)
    pages = sorted(tmp.glob("small-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
    return previews, [Image.open(p).convert("RGB") for p in pages]


def veil_data_url(img: Image.Image) -> str:
    h = max(1, round(VEIL_WIDTH * img.height / img.width))
    tiny = img.resize((VEIL_WIDTH, h), Image.Resampling.BOX)
    buf = io.BytesIO()
    tiny.save(buf, "PNG", optimize=True)
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--slug", required=True, help="URL 조각(소문자·숫자·하이픈), 예: battleview-3d")
    ap.add_argument("--preview", default="1,2,3", help="공개할 슬라이드 번호(1부터), 기본 1,2,3")
    ap.add_argument("--width", type=int, default=1920, help="미리보기 이미지 폭(px)")
    args = ap.parse_args()

    if not SLUG_RE.match(args.slug):
        sys.exit("--slug 는 소문자·숫자·하이픈만 씁니다(예: battleview-3d).")
    prs = Presentation(str(args.pptx))
    count = len(prs.slides)
    preview = parse_preview(args.preview, count)

    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        blobs = full_bleed_images(prs, preview)
        if blobs:
            images = [Image.open(io.BytesIO(b)).convert("RGB") for b in blobs]
            previews = {n: images[n - 1] for n in preview}
            for n, img in previews.items():
                if img.width > args.width:
                    previews[n] = img.resize((args.width, round(args.width * img.height / img.width)), Image.Resampling.LANCZOS)
            smalls = images
            source = "embedded-images"
        else:
            warn_missing_fonts(deck_fonts(prs))
            previews, smalls = render_with_libreoffice(args.pptx, tmp, preview, args.width)
            source = "libreoffice"
        if len(smalls) != count:
            sys.exit(f"렌더된 장수({len(smalls)})가 슬라이드 수({count})와 다릅니다.")

        out_dir = ROOT / "public" / "gallery" / args.slug
        if out_dir.exists():
            shutil.rmtree(out_dir)
        out_dir.mkdir(parents=True)
        entries = []
        for n in preview:
            img = previews[n]
            rel = f"/gallery/{args.slug}/slide-{n:02d}.webp"
            img.save(ROOT / "public" / rel.lstrip("/"), "WEBP", quality=86, method=6)
            entries.append({"slide": n, "src": rel, "width": img.width, "height": img.height})

        veils = [{"slide": i + 1, "dataUrl": veil_data_url(img)} for i, img in enumerate(smalls) if (i + 1) not in preview]

    ratio = [prs.slide_width, prs.slide_height]
    manifest = {
        "slug": args.slug,
        "slideCount": count,
        "aspectRatio": ratio,
        "source": source,
        "previews": entries,
        "veils": veils,
    }
    manifest_path = ROOT / "src" / "content" / "gallery" / f"{args.slug}.deck.json"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    print(f"{args.slug}: {count}장 중 {len(entries)}장 공개 ({source})")
    print(f"  {out_dir.relative_to(ROOT)}/")
    print(f"  {manifest_path.relative_to(ROOT)}")
    print("다음: src/content/gallery.ts 에 작품 정보(제목·용도·수상 등)를 추가하세요.")


if __name__ == "__main__":
    main()
